"""
v2t.exporters - 多格式导出与转录成果排版模块
负责 SRT 字幕生成(带折行优化 Item A5)、排版精美 Word 文档生成、转录统计 JSON (Item E2) 导出。
"""
import os
import re
import json
import logging
from datetime import datetime
from v2t.text_processor import clean_and_improve_text, ms_to_srt_time, wrap_subtitle_text

logger = logging.getLogger("ParaScribe")


def get_audio_info(media_path):
    """获取媒体文件大小与基础信息"""
    try:
        file_size = os.path.getsize(media_path) / (1024 * 1024)
        return {
            "size_mb": round(file_size, 2),
            "filename": os.path.basename(media_path)
        }
    except Exception as e:
        logger.error(f"获取媒体信息失败: {e}")
        return {"size_mb": 0, "filename": str(media_path)}


def export_srt(sentence_info, srt_path, config=None, replacements=None):
    """
    导出标准 SRT 字幕文件。
    集成口语废词过滤、标点净化与行长度智能折行 (Item A5)。
    """
    if not sentence_info:
        return False

    cfg = config or {}
    lines = []
    idx = 1
    for item in sentence_info:
        raw_text = item.get("text", "").strip()
        if not raw_text:
            continue
        cleaned_text = clean_and_improve_text(raw_text, cfg, replacements=replacements)
        if not cleaned_text or not re.search(r'[\u4e00-\u9fa5a-zA-Z0-9]', cleaned_text):
            continue

        # 字幕末尾去除冗余句号，符合主流字幕视觉规范
        sub_text = re.sub(r'[。.]+$', '', cleaned_text).strip()
        if not sub_text:
            continue

        # 智能控制字幕单行字数与折行 (Item A5)
        sub_text = wrap_subtitle_text(sub_text, max_chars_per_line=21)

        start_time = ms_to_srt_time(item.get("start", 0))
        end_time = ms_to_srt_time(item.get("end", item.get("start", 0) + 1000))

        lines.append(f"{idx}\n{start_time} --> {end_time}\n{sub_text}\n")
        idx += 1

    if not lines:
        return False

    with open(srt_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return True


def export_enhanced_docx(docx_path, base_name, media_file, media_info, total_dur,
                         final_text, ts_text, sentence_info, processing_time):
    """
    生成排版精良、层次分明的 Word (.docx) 文档。
    包含正文版、时间轴版以及完整的转录统计摘要。
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        doc = Document()

        # 默认字体设置
        doc.styles['Normal'].font.name = '微软雅黑'
        doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        doc.styles['Normal'].font.size = Pt(11)

        # 1. 标题
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(base_name)
        title_run.font.name = '黑体'
        title_run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_p.paragraph_format.space_after = Pt(10)

        # 元数据说明行
        dur_str = f"{int(total_dur // 60)}分{int(total_dur % 60)}秒" if total_dur > 0 else "未知"
        meta_p = doc.add_paragraph()
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_run = meta_p.add_run(
            f"源媒体: {os.path.basename(media_file)} ({media_info['size_mb']} MB) | "
            f"总时长: {dur_str} | "
            f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        )
        meta_run.font.size = Pt(9.5)
        meta_run.font.color.rgb = RGBColor(128, 128, 128)
        meta_p.paragraph_format.space_after = Pt(20)

        # 2. 一、 正文整理版
        h1 = doc.add_paragraph()
        h1_run = h1.add_run("一、 正文整理版")
        h1_run.font.name = '黑体'
        h1_run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        h1_run.font.size = Pt(14)
        h1_run.font.bold = True
        h1_run.font.color.rgb = RGBColor(30, 60, 120)
        h1.paragraph_format.space_before = Pt(12)
        h1.paragraph_format.space_after = Pt(8)

        for para in final_text.split("\n\n"):
            p_str = para.strip()
            if p_str:
                p = doc.add_paragraph(p_str)
                p.paragraph_format.first_line_indent = Pt(22)
                p.paragraph_format.line_spacing = 1.35
                p.paragraph_format.space_after = Pt(6)

        # 3. 二、 精确时间轴对照版（如果有时间戳）
        if ts_text:
            doc.add_page_break()
            h2 = doc.add_paragraph()
            h2_run = h2.add_run("二、 精确时间轴对照版")
            h2_run.font.name = '黑体'
            h2_run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
            h2_run.font.size = Pt(14)
            h2_run.font.bold = True
            h2_run.font.color.rgb = RGBColor(30, 60, 120)
            h2.paragraph_format.space_before = Pt(12)
            h2.paragraph_format.space_after = Pt(8)

            for line in ts_text.split("\n\n"):
                line_str = line.strip()
                if line_str:
                    p = doc.add_paragraph()
                    # 分离时间戳标签与文字
                    m = re.match(r"^(\[[0-9:]+\s*-\s*[0-9:]+\])\s*(.*)$", line_str)
                    if m:
                        ts_tag, content = m.groups()
                        r_ts = p.add_run(ts_tag + "  ")
                        r_ts.font.size = Pt(9.5)
                        r_ts.font.color.rgb = RGBColor(120, 120, 120)
                        r_ts.font.bold = True
                        r_text = p.add_run(content)
                        r_text.font.size = Pt(10.5)
                    else:
                        p.add_run(line_str)
                    p.paragraph_format.line_spacing = 1.25
                    p.paragraph_format.space_after = Pt(4)

        # 4. 三、 转录统计摘要
        doc.add_page_break()
        h3 = doc.add_paragraph()
        h3_run = h3.add_run("三、 转录统计摘要")
        h3_run.font.name = '黑体'
        h3_run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        h3_run.font.size = Pt(14)
        h3_run.font.bold = True
        h3_run.font.color.rgb = RGBColor(30, 60, 120)
        h3.paragraph_format.space_before = Pt(12)
        h3.paragraph_format.space_after = Pt(12)

        total_chars = len(final_text.replace("\n", "").replace(" ", ""))
        speed_factor = (total_dur / max(0.1, processing_time)) if total_dur > 0 else 0

        stats = [
            ("源媒体文件", os.path.basename(media_file)),
            ("媒体大小", f"{media_info['size_mb']} MB"),
            ("媒体总时长", dur_str),
            ("识别字数统计", f"{total_chars} 字符"),
            ("处理耗时", f"{processing_time:.1f} 秒"),
            ("转录加速比", f"{speed_factor:.2f} x 实时" if speed_factor > 0 else "N/A"),
            ("转写系统", "言迹 (ParaScribe) v2.0"),
            ("核心识别模型", "FunASR SeACo-Paraformer Large (离线增强版)"),
            ("语音端点检测", "FSMN-VAD (单段上限 30s)"),
            ("智能标点恢复", "CT-Transformer Large"),
            ("导出完成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ]

        table = doc.add_table(rows=len(stats), cols=2)
        table.autofit = True
        for row_idx, (k, v) in enumerate(stats):
            cell_k = table.cell(row_idx, 0)
            cell_v = table.cell(row_idx, 1)
            cell_k.text = k
            cell_v.text = v
            cell_k.paragraphs[0].runs[0].font.bold = True
            cell_k.paragraphs[0].runs[0].font.size = Pt(10)
            cell_v.paragraphs[0].runs[0].font.size = Pt(10)

        doc.save(docx_path)
        logger.info(f"已同步生成精排 Word 文档: {docx_path}")
        return True
    except Exception as e:
        logger.warning(f"生成 Word 文档异常（不影响文本转录）: {e}")
        return False


def write_summary_json(summary_path, summary_data):
    """输出结构化的转录任务总汇报 JSON (Item E2)"""
    try:
        with open(summary_path, "w", encoding="utf-8") as f:
            json.dump(summary_data, f, ensure_ascii=False, indent=2)
        logger.info(f"已输出结构化统计摘要: {summary_path}")
        return True
    except Exception as e:
        logger.warning(f"写入摘要 JSON 失败: {e}")
        return False
